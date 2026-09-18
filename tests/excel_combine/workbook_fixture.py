from pathlib import Path

from openpyxl import Workbook


class WorkbookFixture:
    """Writes test .xlsx files at test time (data files must never be committed: the repo is public)."""

    @staticmethod
    def write(path: Path, sheets: dict[str, list[list]]) -> Path:
        """`sheets`: tab name -> rows, where row 1 is the header. `None` cells stay empty."""
        workbook = Workbook()
        workbook.remove(workbook.active)
        for sheet_name, rows in sheets.items():
            sheet = workbook.create_sheet(sheet_name)
            for row_number, row in enumerate(rows, start=1):
                for column_number, value in enumerate(row, start=1):
                    if value is not None:
                        sheet.cell(row=row_number, column=column_number, value=value)
        workbook.save(path)
        return path
