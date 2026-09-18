from datetime import datetime

import pandas as pd
import pytest
from openpyxl import load_workbook

from ximenas_tooling.excel_combine import CombineValidationError, run
from tests.excel_combine.workbook_fixture import WorkbookFixture

TAB = "Data"
OUTPUT = "combined.xlsx"


def read_output(folder, sheet_name=TAB) -> pd.DataFrame:
    return pd.read_excel(folder / OUTPUT, sheet_name=sheet_name, engine="openpyxl")


class TestRunGivenTwoMatchingFiles:
    @pytest.fixture
    def folder(self, tmp_path):
        WorkbookFixture.write(
            tmp_path / "b_report.xlsx",
            {"Summary": [["Other"], [1]], " data ": [["clicks", "DATE"], [3, datetime(2026, 2, 1)]]},
        )
        WorkbookFixture.write(
            tmp_path / "a_report.xlsx",
            {TAB: [["Date", "Clicks"], [datetime(2026, 1, 1), 1], [None, None], [datetime(2026, 1, 2), 2]]},
        )
        run(str(tmp_path), TAB)
        return tmp_path

    def test_THEN_output_sheet_is_named_after_tab(self, folder):
        assert load_workbook(folder / OUTPUT).sheetnames == [TAB]

    def test_THEN_files_are_combined_in_name_order_with_source_file(self, folder):
        assert read_output(folder)["Source File"].tolist() == ["a_report.xlsx", "a_report.xlsx", "b_report.xlsx"]

    def test_THEN_headers_come_from_first_file(self, folder):
        assert list(read_output(folder).columns) == ["Date", "Clicks", "Source File"]

    def test_THEN_empty_rows_are_dropped_and_values_kept(self, folder):
        assert read_output(folder)["Clicks"].tolist() == [1, 2, 3]

    def test_THEN_dates_stay_dates(self, folder):
        assert load_workbook(folder / OUTPUT)[TAB]["A2"].value == datetime(2026, 1, 1)

    def test_WHEN_rerun_THEN_previous_output_is_not_ingested(self, folder):
        run(str(folder), TAB)
        assert len(read_output(folder)) == 3


class TestRunGivenColumnsInDifferentOrderCaseAndSpacing:
    @pytest.fixture
    def folder(self, tmp_path):
        WorkbookFixture.write(tmp_path / "a.xlsx", {TAB: [["Campaign Name", "Clicks", "Spend"], ["x", 1, 10.5]]})
        WorkbookFixture.write(tmp_path / "b.xlsx", {TAB: [[" spend ", "CAMPAIGN   name", "clicks"], [20.5, "y", 2]]})
        return tmp_path

    def test_THEN_headers_come_from_first_file(self, folder):
        run(str(folder), TAB)
        assert list(read_output(folder).columns) == ["Campaign Name", "Clicks", "Spend", "Source File"]

    def test_THEN_every_value_lands_under_its_matching_header(self, folder):
        run(str(folder), TAB)
        assert read_output(folder).to_dict("records") == [
            {"Campaign Name": "x", "Clicks": 1, "Spend": 10.5, "Source File": "a.xlsx"},
            {"Campaign Name": "y", "Clicks": 2, "Spend": 20.5, "Source File": "b.xlsx"},
        ]

    def test_THEN_no_column_mismatch_warning_is_printed(self, folder, capsys):
        run(str(folder), TAB)
        assert "COLUMN MISMATCH" not in capsys.readouterr().out


class TestRunGivenOtherTabsDiffer:
    def test_THEN_only_the_extracted_tab_is_compared(self, tmp_path):
        WorkbookFixture.write(tmp_path / "a.xlsx", {TAB: [["X"], [1]], "Notes": [["Foo"], [1]]})
        WorkbookFixture.write(tmp_path / "b.xlsx", {TAB: [["X"], [2]], "Notes": [["Bar", "Baz"], [1, 2]]})
        run(str(tmp_path), TAB)
        assert len(read_output(tmp_path)) == 2


class TestRunGivenColumnMismatch:
    @pytest.fixture
    def folder(self, tmp_path):
        WorkbookFixture.write(tmp_path / "a.xlsx", {TAB: [["Campaign", "Clicks"], ["x", 1]]})
        WorkbookFixture.write(tmp_path / "b.xlsx", {TAB: [["Campaign", "Spend"], ["y", 5]]})
        return tmp_path

    def test_WHEN_default_THEN_error_names_the_override(self, folder):
        with pytest.raises(CombineValidationError, match="allow_column_mismatch"):
            run(str(folder), TAB)

    def test_WHEN_default_THEN_nothing_is_written(self, folder):
        with pytest.raises(CombineValidationError):
            run(str(folder), TAB)
        assert not (folder / OUTPUT).exists()

    def test_WHEN_allowed_THEN_union_with_blanks_is_written(self, folder):
        run(str(folder), TAB, allow_column_mismatch=True)
        assert read_output(folder)["Spend"].isna().tolist() == [True, False]

    def test_WHEN_allowed_THEN_warning_is_printed(self, folder, capsys):
        run(str(folder), TAB, allow_column_mismatch=True)
        assert "WARNING: COLUMN MISMATCH" in capsys.readouterr().out


class TestRunGivenTabProblems:
    def test_GIVEN_tab_case_and_spacing_differ_THEN_tab_is_found(self, tmp_path):
        WorkbookFixture.write(tmp_path / "a.xlsx", {"  Monthly   DATA ": [["X"], [1]]})
        run(str(tmp_path), "monthly data")
        assert len(read_output(tmp_path, sheet_name="monthly data")) == 1

    def test_GIVEN_missing_tab_THEN_error_lists_file_and_its_tabs(self, tmp_path):
        WorkbookFixture.write(tmp_path / "a.xlsx", {"Summary": [["X"], [1]], "Raw": [["X"], [1]]})
        with pytest.raises(CombineValidationError, match=r"a\.xlsx: tab \[Data\] not found\. Tabs in this file: \['Summary', 'Raw'\]"):
            run(str(tmp_path), TAB)

    def test_GIVEN_two_tabs_match_THEN_error(self, tmp_path):
        WorkbookFixture.write(tmp_path / "a.xlsx", {"Data": [["X"], [1]], "DATA ": [["X"], [1]]})
        with pytest.raises(CombineValidationError, match="several tabs match"):
            run(str(tmp_path), TAB)

    def test_GIVEN_duplicate_headers_in_file_THEN_error(self, tmp_path):
        WorkbookFixture.write(tmp_path / "a.xlsx", {TAB: [["Clicks", "Clicks"], [1, 2]]})
        with pytest.raises(CombineValidationError, match="several columns have the same name"):
            run(str(tmp_path), TAB)

    def test_GIVEN_data_under_empty_header_THEN_error_names_column_letter(self, tmp_path):
        WorkbookFixture.write(tmp_path / "a.xlsx", {TAB: [["Clicks", None], [1, 2]]})
        with pytest.raises(CombineValidationError, match=r"column\(s\) \['B'\]"):
            run(str(tmp_path), TAB)

    def test_GIVEN_problems_in_several_files_THEN_all_are_in_one_error(self, tmp_path):
        WorkbookFixture.write(tmp_path / "a.xlsx", {TAB: [["X"], [1]]})
        WorkbookFixture.write(tmp_path / "b.xlsx", {"Other": [["X"], [1]]})
        WorkbookFixture.write(tmp_path / "c.xlsx", {TAB: [["X", "x"], [1, 2]]})
        with pytest.raises(CombineValidationError) as error:
            run(str(tmp_path), TAB)
        assert len(error.value.problems) == 2


class TestRunFileDiscovery:
    def test_GIVEN_lock_file_THEN_it_is_skipped(self, tmp_path):
        WorkbookFixture.write(tmp_path / "a.xlsx", {TAB: [["X"], [1]]})
        (tmp_path / "~$a.xlsx").write_bytes(b"not a workbook")
        run(str(tmp_path), TAB)
        assert len(read_output(tmp_path)) == 1

    def test_GIVEN_subfolder_xlsx_THEN_it_is_ignored(self, tmp_path):
        WorkbookFixture.write(tmp_path / "a.xlsx", {TAB: [["X"], [1]]})
        (tmp_path / "sub").mkdir()
        WorkbookFixture.write(tmp_path / "sub" / "b.xlsx", {TAB: [["X"], [2]]})
        run(str(tmp_path), TAB)
        assert len(read_output(tmp_path)) == 1

    def test_GIVEN_no_input_files_THEN_error_names_folder(self, tmp_path):
        (tmp_path / "notes.csv").write_text("a,b")
        with pytest.raises(CombineValidationError, match=str(tmp_path)):
            run(str(tmp_path), TAB)

    def test_GIVEN_tab_name_longer_than_excel_limit_THEN_output_sheet_name_is_truncated(self, tmp_path):
        # Real tabs are at most 31 chars, but a typed name can be longer through extra spacing.
        typed_tab = "Campaign    performance    by    week"
        WorkbookFixture.write(tmp_path / "a.xlsx", {"Campaign performance by week": [["X"], [1]]})
        run(str(tmp_path), typed_tab)
        assert load_workbook(tmp_path / OUTPUT).sheetnames == [typed_tab[:31]]
